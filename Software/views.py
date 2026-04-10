from datetime import datetime
from datetime import date, timedelta
from decimal import Decimal
import csv
import json
import io
import math
import os
import shutil
import tempfile
import zipfile
import sqlite3
from pathlib import Path
from urllib.parse import urlencode, urlparse, parse_qsl, urlunparse
from django.conf import settings
from django.contrib import messages
from django.contrib.auth import logout as auth_logout
from django.db import connections
from django.http import JsonResponse , HttpResponse
from django.forms import formset_factory , modelformset_factory
from django.shortcuts import render , HttpResponseRedirect , get_object_or_404
from django.db import transaction
from django.utils import timezone
from django.db.models import Sum, Count, F, Q, DecimalField, ExpressionWrapper
from django.db.models.functions import Coalesce, TruncMonth
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
import Software.models as md


#####################################################################################################################
##		SECONDARY Functions
#####################################################################################################################
def api(request , model=None , id=None):
	if (model == 'client') :
		try:
			obj = md.Client.objects.get(id=id)
			return JsonResponse({"name" : obj.name , "address" : obj.address , "mobile" : obj.mobile , "gstin" : obj.gstin , "state" : obj.state, "debit" : obj.current_debit_balance()})
		except:
			return JsonResponse({"you get what you deserve" : "nothing"})
	elif (model == 'invoices') :
		try:
			obj = md.Bill.objects.get(id=id, is_deleted=False)
			before = md.audit_snapshot(obj, include_related=True)
			obj.is_cleared = not obj.is_cleared
			obj.save()
			md.log_audit_entry(
				"toggle_clear",
				obj,
				actor=request.user,
				before_data=before,
				after_data=md.audit_snapshot(obj, include_related=True),
			)
			return JsonResponse({})
		except:
			pass
	elif (model == 'payments') :
		try:
			obj = md.Payment.objects.get(id=id, is_deleted=False)
			before = md.audit_snapshot(obj)
			obj.is_cleared = not obj.is_cleared
			obj.save()
			md.log_audit_entry(
				"toggle_clear",
				obj,
				actor=request.user,
				before_data=before,
				after_data=md.audit_snapshot(obj),
			)
			return JsonResponse({})
		except:
			pass
	elif (model == 'product') :
		try:
			id = str(id)
			if md.Product.objects.filter(name=id).exists():
				product = md.Product.objects.get(name=id)
				if not product.hsn_code:
					product.hsn_code = md.DEFAULT_HSN_CODE
					product.save(update_fields=["hsn_code"])
			else:
				product = md.Product(name=str(id),hsn_code=md.DEFAULT_HSN_CODE)
				product.save()
			return JsonResponse({"id" : product.id })
		except:
			return JsonResponse({"id":"0"})
	elif (model == 'size') :
		try:
			id = str(id)
			if md.Size.objects.filter(name=id).exists():
				size = md.Size.objects.get(name=id)
			else:
				size = md.Size(name=str(id))
				size.save()
			return JsonResponse({"id" : size.id })
		except:
			return JsonResponse({"id":"0"})
	return None


def logout_view(request):
	auth_logout(request)
	return HttpResponseRedirect("/login/")

def get_invoice_number(parcel_digits=None):
	return md.next_invoice_number(parcel_digits)


def get_firm():
	return md.FirmSettings.current()


def _firm_contact_numbers(firm):
	return [number for number in [firm.contact_number, firm.alternate_contact_number, getattr(firm, "third_contact_number", None)] if number]


def _invoice_terms_lines(firm):
	raw_terms = (getattr(firm, "invoice_terms", "") or "").strip()
	if not raw_terms:
		return [
			"Goods once sold cannot be taken back",
			"Subject to Ahmedabad Jurisdiction",
			"After due date 24% interest will be charged",
		]
	lines = []
	for line in raw_terms.splitlines():
		clean = line.strip()
		while clean and clean[0] in {"✍", "•", "●", "-", "*", "–", "—", "☛", "►"}:
			clean = clean[1:].strip()
		if clean:
			lines.append(clean)
	return lines or [
		"Goods once sold cannot be taken back",
		"Subject to Ahmedabad Jurisdiction",
		"After due date 24% interest will be charged",
	]


def _invoice_template_name():
	firm = get_firm()
	if firm.invoice_template == md.INVOICE_TEMPLATE_INVOICE1:
		return "invoice1.html"
	return "single_invoice.html"


def _print_invoice_template_name():
	firm = get_firm()
	if firm.invoice_template == md.INVOICE_TEMPLATE_INVOICE1:
		return "invoice1.html"
	return "print_invoice.html"


def _invoice_pages(invoice, rows_per_page=16):
	items = list(invoice.item_set.select_related("product", "size").all())
	total_pages = max(1, math.ceil(len(items) / rows_per_page))
	pages = []
	for page_index in range(total_pages):
		start = page_index * rows_per_page
		chunk = items[start:start + rows_per_page]
		pages.append({
			"rows": [
				{"serial": start + offset + 1, "item": item}
				for offset, item in enumerate(chunk)
			],
			"blank_rows": range(rows_per_page - len(chunk)),
			"show_continued": page_index < total_pages - 1,
			"is_last": page_index == total_pages - 1,
			"page_index": page_index,
			"total_pages": total_pages,
		})
	return pages


def _firm_tax_rates():
	firm = get_firm()
	cgst_rate = firm.cgst_rate if firm.cgst_rate is not None else md.DEFAULT_CGST_RATE
	sgst_rate = firm.sgst_rate if firm.sgst_rate is not None else md.DEFAULT_SGST_RATE
	igst_rate = firm.igst_rate if firm.igst_rate is not None else md.DEFAULT_IGST_RATE
	return firm, cgst_rate, sgst_rate, igst_rate


def _apply_bill_tax_rates(bill):
	_, cgst_rate, sgst_rate, igst_rate = _firm_tax_rates()
	if bill.is_taxable:
		if bill.client.state_code == "24":
			bill.cgst = cgst_rate
			bill.sgst = sgst_rate
			bill.igst = Decimal("0.00")
		else:
			bill.cgst = Decimal("0.00")
			bill.sgst = Decimal("0.00")
			bill.igst = igst_rate
	else:
		bill.cgst = Decimal("0.00")
		bill.sgst = Decimal("0.00")
		bill.igst = Decimal("0.00")


def _redirect_with_client(next_url, client_id):
	if not next_url:
		return None
	parsed = urlparse(next_url)
	if parsed.scheme or parsed.netloc:
		return None
	query_params = dict(parse_qsl(parsed.query))
	query_params["client"] = str(client_id)
	new_query = urlencode(query_params)
	return urlunparse((parsed.scheme, parsed.netloc, parsed.path, parsed.params, new_query, parsed.fragment))

def process_agent(agent):
	if agent == '' : return ''
	if agent.isdigit()  and md.Agent.objects.filter(id=int(agent)).exists() : return agent
	new_agent = md.Agent(name=agent , mobile='')
	new_agent.save()
	return str(new_agent.id)


def _refresh_client_balance(client):
	if client is None:
		return Decimal("0.00")
	return client.refresh_debit_balance()


def _refresh_bill_totals(bill):
	if bill is None:
		return None
	bill.qty_total = int(sum((item.qty for item in bill.item_set.all()), Decimal("0")))
	bill.total_before_tax = Decimal("0.00")
	for item in bill.item_set.all():
		bill.total_before_tax += item.rate * item.qty
	bill.total_before_tax *= (Decimal("100.00") - bill.discount) * Decimal("0.01")
	_apply_bill_tax_rates(bill)
	if bill.client.state_code == "24":
		tax_multiplier = Decimal("1.00") + ((bill.cgst + bill.sgst) / Decimal("100.00"))
	else:
		tax_multiplier = Decimal("1.00") + (bill.igst / Decimal("100.00"))
	bill.total_after_tax = Decimal(math.ceil(float(bill.total_before_tax) * float(tax_multiplier)))
	bill.save(update_fields=["qty_total", "total_before_tax", "cgst", "sgst", "igst", "total_after_tax"])
	return bill

def _month_start_range(start_value, end_value):
	current = date(start_value.year, start_value.month, 1)
	while current <= end_value:
		yield current
		if current.month == 12:
			current = date(current.year + 1, 1, 1)
		else:
			current = date(current.year, current.month + 1, 1)

def _format_export_value(value):
	if isinstance(value, Decimal):
		return format(value, "f")
	return value

def _invoice_tax_amount():
	return ExpressionWrapper(F("total_after_tax") - F("total_before_tax"), output_field=DecimalField(max_digits=14, decimal_places=2))

def _tax_totals_for_bills(bills):
	domestic_tax_total = bills.filter(is_taxable=True, cgst__gt=0).aggregate(
		total=Coalesce(Sum(_invoice_tax_amount()), Decimal("0.00"))
	)["total"]
	igst_total = bills.filter(is_taxable=True, igst__gt=0).aggregate(
		total=Coalesce(Sum(_invoice_tax_amount()), Decimal("0.00"))
	)["total"]
	return {
		"cgst_total": domestic_tax_total / 2,
		"sgst_total": domestic_tax_total / 2,
		"igst_total": igst_total,
		"cgst_sgst_total": domestic_tax_total,
		"total_tax": domestic_tax_total + igst_total,
	}

def _build_reports_payload():
	bills = md.Bill.objects.filter(is_deleted=False)
	payments = md.Payment.objects.filter(is_deleted=False)
	invoice_count = bills.count()
	payment_count = payments.count()
	today = timezone.localdate()

	def _month_bounds():
		date_candidates = []
		first_bill = bills.order_by("invoice_date").first()
		last_bill = bills.order_by("-invoice_date").first()
		first_payment = payments.order_by("payment_date").first()
		last_payment = payments.order_by("-payment_date").first()
		for obj, field in ((first_bill, "invoice_date"), (last_bill, "invoice_date"), (first_payment, "payment_date"), (last_payment, "payment_date")):
			if obj is not None:
				date_candidates.append(getattr(obj, field))
		if not date_candidates:
			start = today.replace(day=1)
			return start, today
		start = min(date_candidates).replace(day=1)
		end = max(date_candidates)
		return start, end

	start_date, end_date = _month_bounds()

	sales_total = bills.aggregate(total=Coalesce(Sum("total_after_tax"), Decimal("0.00")))["total"]
	collections_total = payments.aggregate(total=Coalesce(Sum("paid_amount"), Decimal("0.00")))["total"]
	pending_invoices_total = bills.filter(is_cleared=False).aggregate(total=Coalesce(Sum("total_after_tax"), Decimal("0.00")))["total"]
	gst_summary = bills.filter(is_taxable=True).aggregate(
		base_total=Coalesce(Sum("total_before_tax"), Decimal("0.00")),
		taxable_total=Coalesce(Sum("total_after_tax"), Decimal("0.00")),
	)
	gst_summary.update(_tax_totals_for_bills(bills))

	sales_by_month = {
		row["month"]: row["total"]
		for row in bills.annotate(month=TruncMonth("invoice_date")).values("month").annotate(total=Coalesce(Sum("total_after_tax"), Decimal("0.00")))
	}
	collections_by_month = {
		row["month"]: row["total"]
		for row in payments.annotate(month=TruncMonth("payment_date")).values("month").annotate(total=Coalesce(Sum("paid_amount"), Decimal("0.00")))
	}

	monthly_rows = []
	for month in _month_start_range(start_date, end_date):
		sales = sales_by_month.get(month, Decimal("0.00"))
		collections = collections_by_month.get(month, Decimal("0.00"))
		monthly_rows.append({
			"month": month,
			"sales": sales,
			"collections": collections,
			"net_movement": sales - collections,
		})

	client_due_map = {}
	for row in bills.values("client_id", "client__name").annotate(invoice_total=Coalesce(Sum("total_after_tax"), Decimal("0.00")), invoice_count=Count("id")).order_by("client__name"):
		client_due_map[row["client_id"]] = {
			"client_id": row["client_id"],
			"name": row["client__name"],
			"opening_balance": Decimal("0.00"),
			"invoice_total": row["invoice_total"],
			"payment_total": Decimal("0.00"),
			"closing_balance": row["invoice_total"],
			"invoice_count": row["invoice_count"],
		}
	for row in payments.values("client_id").annotate(payment_total=Coalesce(Sum("paid_amount"), Decimal("0.00"))):
		client_data = client_due_map.get(row["client_id"])
		if client_data:
			client_data["payment_total"] = row["payment_total"]
			client_data["closing_balance"] = client_data["invoice_total"] - row["payment_total"]
		else:
			client = md.Client.objects.filter(id=row["client_id"]).first()
			client_due_map[row["client_id"]] = {
				"client_id": row["client_id"],
				"name": client.name if client else "Unknown",
				"opening_balance": Decimal("0.00"),
				"invoice_total": Decimal("0.00"),
				"payment_total": row["payment_total"],
				"closing_balance": -row["payment_total"],
				"invoice_count": 0,
			}
	client_due_rows = sorted([row for row in client_due_map.values() if row["closing_balance"] > 0], key=lambda x: x["closing_balance"], reverse=True)
	total_receivables = sum((row["closing_balance"] for row in client_due_rows), Decimal("0.00"))

	aging_source = list(bills.filter(is_cleared=False).order_by("invoice_date"))
	aging_buckets = [
		{"label": "0-30 Days", "count": 0, "amount": Decimal("0.00")},
		{"label": "31-60 Days", "count": 0, "amount": Decimal("0.00")},
		{"label": "61-90 Days", "count": 0, "amount": Decimal("0.00")},
		{"label": "90+ Days", "count": 0, "amount": Decimal("0.00")},
	]
	for bill in aging_source:
		age_days = max((today - bill.invoice_date).days, 0)
		if age_days <= 30:
			target = aging_buckets[0]
		elif age_days <= 60:
			target = aging_buckets[1]
		elif age_days <= 90:
			target = aging_buckets[2]
		else:
			target = aging_buckets[3]
		target["count"] += 1
		target["amount"] += bill.total_after_tax

	agent_rows = (
		bills.exclude(agent__isnull=True)
		.values("agent__id", "agent__name")
		.annotate(
			invoice_count=Count("id"),
			total=Coalesce(Sum("total_after_tax"), Decimal("0.00")),
		)
		.order_by("-total", "agent__name")
	)

	product_amount = ExpressionWrapper(F("qty") * F("rate"), output_field=DecimalField(max_digits=14, decimal_places=2))
	product_rows = (
		md.Item.objects.filter(bill__is_deleted=False)
		.values("product__id", "product__name")
		.annotate(
			qty_total=Coalesce(Sum("qty"), Decimal("0.00")),
			total=Coalesce(Sum(product_amount), Decimal("0.00")),
			line_count=Count("id"),
		)
		.order_by("-total", "product__name")
	)

	recent_activity = []
	for bill in bills.order_by("-invoice_date", "-id").select_related("client", "agent")[:10]:
		recent_activity.append({
			"kind": "Invoice",
			"date": bill.invoice_date,
			"title": bill.display_invoice_number(),
			"party": bill.client.name,
			"amount": bill.total_after_tax,
			"status": "Cleared" if bill.is_cleared else "Pending",
			"url": "/invoices/{0}".format(bill.id),
		})
	for payment in payments.order_by("-payment_date", "-id").select_related("client")[:10]:
		recent_activity.append({
			"kind": "Payment",
			"date": payment.payment_date,
			"title": payment.method_of_payment,
			"party": payment.client.name,
			"amount": payment.paid_amount,
			"status": "Cleared" if payment.is_cleared else "Pending",
			"url": "/payments/{0}".format(payment.id),
		})
	recent_activity = sorted(recent_activity, key=lambda row: (row["date"], row["title"]), reverse=True)[:10]

	top_products = (
		md.Item.objects.filter(bill__is_deleted=False)
		.values("product__id", "product__name")
		.annotate(
			total_qty=Coalesce(Sum("qty"), Decimal("0.00")),
			total_amount=Coalesce(Sum(product_amount), Decimal("0.00")),
			line_count=Count("id"),
		)
		.order_by("-total_amount", "product__name")[:5]
	)
	top_agents = (
		bills.exclude(agent__isnull=True)
		.values("agent__id", "agent__name")
		.annotate(
			total_amount=Coalesce(Sum("total_after_tax"), Decimal("0.00")),
			invoice_count=Count("id"),
		)
		.order_by("-total_amount", "agent__name")[:5]
	)

	gst_cards = [
		{"label": "Sales", "value": gst_summary["taxable_total"], "subtext": "All records"},
		{"label": "CGST + SGST", "value": gst_summary["cgst_sgst_total"], "subtext": "Domestic invoices"},
		{"label": "IGST", "value": gst_summary["igst_total"], "subtext": "Interstate invoices"},
	]
	return {
		"start_date": start_date,
		"end_date": end_date,
		"bills": bills,
		"payments": payments,
		"sales_total": sales_total,
		"collections_total": collections_total,
		"pending_invoices_total": pending_invoices_total,
		"gst_summary": gst_summary,
		"monthly_rows": monthly_rows,
		"receivables": client_due_rows,
		"total_receivables": total_receivables,
		"aging_buckets": aging_buckets,
		"agent_rows": agent_rows,
		"product_rows": product_rows,
		"pending_invoice_count": len(aging_source),
		"invoice_count": invoice_count,
		"payment_count": payment_count,
		"recent_activity": recent_activity,
		"top_products": top_products,
		"top_agents": top_agents,
		"gst_cards": gst_cards,
	}

def _style_workbook_sheet(ws):
	for cell in ws[1]:
		cell.font = Font(bold=True)
		cell.alignment = Alignment(horizontal="center")
	ws.freeze_panes = "A2"
	ws.auto_filter.ref = ws.dimensions

def _append_sheet_rows(ws, headers, rows):
	ws.append(headers)
	for row in rows:
		ws.append([_format_export_value(value) for value in row])
	_style_workbook_sheet(ws)

def _csv_bytes(headers, rows):
	buffer = io.StringIO()
	writer = csv.writer(buffer)
	writer.writerow(headers)
	for row in rows:
		writer.writerow([_format_export_value(value) for value in row])
	return buffer.getvalue().encode("utf-8-sig")

def _reports_filename(extension):
	return f"reports.{extension}"

def _backup_folder():
	backup_dir = Path(settings.BASE_DIR) / "backups"
	backup_dir.mkdir(parents=True, exist_ok=True)
	return backup_dir

def _build_backup_archive():
	buffer = io.BytesIO()
	manifest = {
		"generated_at": timezone.now().isoformat(),
		"project": "AccountManager",
		"database": os.path.basename(settings.DATABASES["default"]["NAME"]),
		"media_root": os.path.basename(settings.MEDIA_ROOT),
		"firm_name": str(get_firm()),
	}
	db_path = Path(settings.DATABASES["default"]["NAME"])
	media_root = Path(settings.MEDIA_ROOT)
	with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
		archive.writestr("manifest.json", json.dumps(manifest, indent=2, default=str))
		if db_path.exists():
			archive.write(db_path, arcname="database/db.sqlite3")
		if media_root.exists():
			for file_path in media_root.rglob("*"):
				if file_path.is_file():
					archive.write(file_path, arcname=f"media/{file_path.relative_to(media_root)}")
	buffer.seek(0)
	return buffer

def _pick_backup_db_member(names):
	preferred = [
		"database/db.sqlite3",
		"db.sqlite3",
	]
	for candidate in preferred:
		if candidate in names:
			return candidate
	for name in names:
		lowered = name.lower()
		if lowered.endswith(".sqlite3") or lowered.endswith(".db"):
			return name
	return None

def _save_backup_snapshot(prefix="pre_restore"):
	backup_dir = _backup_folder()
	timestamp = timezone.localtime().strftime("%Y%m%d_%H%M%S")
	snapshot_path = backup_dir / f"{prefix}_{timestamp}.zip"
	snapshot_path.write_bytes(_build_backup_archive().getvalue())
	return snapshot_path

def _restore_from_backup_file(uploaded_file):
	with tempfile.TemporaryDirectory() as temp_dir:
		temp_dir_path = Path(temp_dir)
		upload_path = temp_dir_path / "upload.zip"
		with open(upload_path, "wb") as handle:
			for chunk in uploaded_file.chunks():
				handle.write(chunk)

		extract_dir = temp_dir_path / "extract"
		extract_dir.mkdir(parents=True, exist_ok=True)
		with zipfile.ZipFile(upload_path, "r") as archive:
			names = archive.namelist()
			db_member = _pick_backup_db_member(names)
			if db_member is None:
				raise ValueError("Backup file does not contain a SQLite database. Expected database/db.sqlite3 or db.sqlite3.")
			archive.extractall(extract_dir)

		db_source = extract_dir / Path(db_member)
		media_source = extract_dir / "media"
		if not media_source.exists():
			media_source = extract_dir / "media".replace("/", os.sep)
		db_target = Path(settings.DATABASES["default"]["NAME"])
		media_target = Path(settings.MEDIA_ROOT)

		snapshot_path = _save_backup_snapshot("pre_restore")
		connections.close_all()

		source_conn = sqlite3.connect(str(db_source))
		target_conn = sqlite3.connect(str(db_target))
		try:
			source_conn.backup(target_conn)
			target_conn.commit()
		finally:
			target_conn.close()
			source_conn.close()

		if media_source.exists():
			if media_target.exists():
				shutil.rmtree(media_target)
			shutil.copytree(media_source, media_target)

		connections.close_all()
		return snapshot_path

def _truncate_company_database():
	snapshot_path = _save_backup_snapshot("pre_truncate")
	models_to_clear = [
		md.Item,
		md.Bill,
		md.Payment,
		md.InvoiceSequence,
		md.AuditLog,
		md.Product,
	]
	with transaction.atomic():
		with connections["default"].constraint_checks_disabled():
			for model in models_to_clear:
				model.objects.all().delete()
			md.Client.objects.all().update(debit_balance=Decimal("0.00"))
	return snapshot_path

#####################################################################################################################
##         SECTION PAGES      
#####################################################################################################################	
def homepage(request):
	template="homepage.html"
	today = timezone.localdate()
	movement_sort = request.GET.get("movement_sort", "asc").lower()
	bills = md.Bill.objects.filter(is_deleted=False)
	payments = md.Payment.objects.filter(is_deleted=False)
	month_start = today.replace(day=1)

	def _month_bounds():
		date_candidates = []
		first_bill = bills.order_by("invoice_date").first()
		last_bill = bills.order_by("-invoice_date").first()
		first_payment = payments.order_by("payment_date").first()
		last_payment = payments.order_by("-payment_date").first()
		for obj, field in ((first_bill, "invoice_date"), (last_bill, "invoice_date"), (first_payment, "payment_date"), (last_payment, "payment_date")):
			if obj is not None:
				date_candidates.append(getattr(obj, field))
		if not date_candidates:
			start = today.replace(day=1)
			return start, today
		return min(date_candidates).replace(day=1), max(date_candidates)

	start_date, end_date = _month_bounds()
	month_start = max(month_start, start_date)

	sales_total = bills.aggregate(total=Coalesce(Sum("total_after_tax"), Decimal("0.00")))["total"]
	collections_total = payments.aggregate(total=Coalesce(Sum("paid_amount"), Decimal("0.00")))["total"]
	invoice_count = bills.count()
	payment_count = payments.count()
	pending_bills = bills.filter(is_cleared=False)
	cleared_bills = bills.filter(is_cleared=True)
	pending_invoice_count = pending_bills.count()
	cleared_invoice_count = cleared_bills.count()
	pending_invoice_total = pending_bills.aggregate(total=Coalesce(Sum("total_after_tax"), Decimal("0.00")))["total"]
	overdue_total = pending_bills.filter(invoice_date__lt=today.replace(day=1)).aggregate(total=Coalesce(Sum("total_after_tax"), Decimal("0.00")))["total"]
	month_bills = bills.filter(invoice_date__gte=month_start)
	month_payments = payments.filter(payment_date__gte=month_start)
	month_sales = month_bills.aggregate(total=Coalesce(Sum("total_after_tax"), Decimal("0.00")))["total"]
	month_collections = month_payments.aggregate(total=Coalesce(Sum("paid_amount"), Decimal("0.00")))["total"]
	prev_month_end = month_start - timedelta(days=1)
	prev_month_start = prev_month_end.replace(day=1)
	prev_month_sales = bills.filter(invoice_date__range=(prev_month_start, prev_month_end)).aggregate(total=Coalesce(Sum("total_after_tax"), Decimal("0.00")))["total"]
	prev_month_collections = payments.filter(payment_date__range=(prev_month_start, prev_month_end)).aggregate(total=Coalesce(Sum("paid_amount"), Decimal("0.00")))["total"]

	client_due_map = {}
	for row in bills.values("client_id", "client__name").annotate(invoice_total=Coalesce(Sum("total_after_tax"), Decimal("0.00")), invoice_count=Count("id")).order_by("client__name"):
		client_due_map[row["client_id"]] = {
			"client_id": row["client_id"],
			"name": row["client__name"],
			"opening_balance": Decimal("0.00"),
			"invoice_total": row["invoice_total"],
			"payment_total": Decimal("0.00"),
			"closing_balance": row["invoice_total"],
			"invoice_count": row["invoice_count"],
		}
	for row in payments.values("client_id").annotate(payment_total=Coalesce(Sum("paid_amount"), Decimal("0.00"))):
		client_data = client_due_map.get(row["client_id"])
		if client_data:
			client_data["payment_total"] = row["payment_total"]
			client_data["closing_balance"] = client_data["invoice_total"] - row["payment_total"]
		else:
			client = md.Client.objects.filter(id=row["client_id"]).first()
			client_due_map[row["client_id"]] = {
				"client_id": row["client_id"],
				"name": client.name if client else "Unknown",
				"opening_balance": Decimal("0.00"),
				"invoice_total": Decimal("0.00"),
				"payment_total": row["payment_total"],
				"closing_balance": -row["payment_total"],
				"invoice_count": 0,
			}
	client_due_rows = sorted([row for row in client_due_map.values() if row["closing_balance"] > 0], key=lambda x: x["closing_balance"], reverse=True)
	total_receivables = sum((row["closing_balance"] for row in client_due_rows), Decimal("0.00"))

	aging_source = list(pending_bills.order_by("invoice_date"))
	aging_buckets = [
		{"label": "0-30 Days", "count": 0, "amount": Decimal("0.00")},
		{"label": "31-60 Days", "count": 0, "amount": Decimal("0.00")},
		{"label": "61-90 Days", "count": 0, "amount": Decimal("0.00")},
		{"label": "90+ Days", "count": 0, "amount": Decimal("0.00")},
	]
	for bill in aging_source:
		age_days = max((today - bill.invoice_date).days, 0)
		if age_days <= 30:
			target = aging_buckets[0]
		elif age_days <= 60:
			target = aging_buckets[1]
		elif age_days <= 90:
			target = aging_buckets[2]
		else:
			target = aging_buckets[3]
		target["count"] += 1
		target["amount"] += bill.total_after_tax

	def next_month(value):
		if value.month == 12:
			return date(value.year + 1, 1, 1)
		return date(value.year, value.month + 1, 1)

	def month_start_range(start_value, end_value):
		current = date(start_value.year, start_value.month, 1)
		while current <= end_value:
			yield current
			current = next_month(current)

	sales_by_month = {}
	for row in bills.annotate(month=TruncMonth("invoice_date")).values("month").annotate(total=Coalesce(Sum("total_after_tax"), Decimal("0.00"))):
		month_key = row["month"].date() if hasattr(row["month"], "date") else row["month"]
		sales_by_month[month_key] = row["total"]
	collections_by_month = {}
	for row in payments.annotate(month=TruncMonth("payment_date")).values("month").annotate(total=Coalesce(Sum("paid_amount"), Decimal("0.00"))):
		month_key = row["month"].date() if hasattr(row["month"], "date") else row["month"]
		collections_by_month[month_key] = row["total"]
	monthly_rows = []
	for month in month_start_range(start_date, end_date):
		sales = sales_by_month.get(month, Decimal("0.00"))
		collections = collections_by_month.get(month, Decimal("0.00"))
		monthly_rows.append({
			"month": month,
			"sales": sales,
			"collections": collections,
			"net_movement": sales - collections,
		})
	if movement_sort == "desc":
		monthly_rows.reverse()

	taxable_total = bills.filter(is_taxable=True).aggregate(total=Coalesce(Sum("total_after_tax"), Decimal("0.00")))["total"]
	tax_totals = _tax_totals_for_bills(bills)
	gst_cards = [
		{"label": "Sales", "value": taxable_total, "subtext": "All records"},
		{"label": "CGST + SGST", "value": tax_totals["cgst_sgst_total"], "subtext": "Domestic invoices"},
		{"label": "IGST", "value": tax_totals["igst_total"], "subtext": "Interstate invoices"},
	]

	recent_activity = []
	for bill in bills.order_by("-invoice_date", "-id").select_related("client", "agent")[:10]:
		recent_activity.append({
			"kind": "Invoice",
			"date": bill.invoice_date,
			"title": bill.display_invoice_number(),
			"party": bill.client.name,
			"amount": bill.total_after_tax,
			"status": "Cleared" if bill.is_cleared else "Pending",
			"url": "/invoices/{0}".format(bill.id),
		})
	for payment in payments.order_by("-payment_date", "-id").select_related("client")[:10]:
		recent_activity.append({
			"kind": "Payment",
			"date": payment.payment_date,
			"title": payment.method_of_payment,
			"party": payment.client.name,
			"amount": payment.paid_amount,
			"status": "Cleared" if payment.is_cleared else "Pending",
			"url": "/payments/{0}".format(payment.id),
		})
	recent_activity = sorted(recent_activity, key=lambda row: (row["date"], row["title"]), reverse=True)[:10]

	product_amount = ExpressionWrapper(F("qty") * F("rate"), output_field=DecimalField(max_digits=14, decimal_places=2))
	top_products = (
		md.Item.objects.filter(bill__is_deleted=False)
		.values("product__id", "product__name")
		.annotate(
			total_qty=Coalesce(Sum("qty"), Decimal("0.00")),
			total_amount=Coalesce(Sum(product_amount), Decimal("0.00")),
			line_count=Count("id"),
		)
		.order_by("-total_amount", "product__name")[:5]
	)
	top_agents = (
		bills.exclude(agent__isnull=True)
		.values("agent__id", "agent__name")
		.annotate(
			total_amount=Coalesce(Sum("total_after_tax"), Decimal("0.00")),
			invoice_count=Count("id"),
		)
		.order_by("-total_amount", "agent__name")[:5]
	)

	context = {
		"current_date": today,
		"overview_label": "All Records",
		"client_count": md.Client.objects.count(),
		"invoice_count": invoice_count,
		"payment_count": payment_count,
		"pending_invoice_count": pending_invoice_count,
		"cleared_invoice_count": cleared_invoice_count,
		"sales_total": sales_total,
		"collections_total": collections_total,
		"pending_invoice_total": pending_invoice_total,
		"overdue_total": overdue_total,
		"month_sales": month_sales,
		"month_collections": month_collections,
		"last_month_sales": prev_month_sales,
		"last_month_collections": prev_month_collections,
		"month_sales_change": month_sales - prev_month_sales,
		"month_collections_change": month_collections - prev_month_collections,
		"total_receivables": total_receivables,
		"aging_buckets": aging_buckets,
		"monthly_rows": monthly_rows,
		"movement_sort": movement_sort,
		"top_receivables": client_due_rows[:5],
		"gst_cards": gst_cards,
		"total_tax": tax_totals["total_tax"],
		"recent_activity": recent_activity,
		"top_products": top_products,
		"top_agents": top_agents,
	}
	return render(request,template , context)

def agents(request):
	agents = md.Agent.objects.all()
	template="agents.html"
	context = {'agents' : agents, 'firm': get_firm()}
	return render(request,template , context)

def clients(request):
	clients = md.Client.objects.all()
	template="clients.html"
	context = {'clients' : clients}
	return render(request,template , context)

def invoices(request):
	invoices = md.Bill.objects.filter(is_deleted=False)
	template="invoices.html"
	context = {'invoices' : invoices}
	return render(request,template , context)

def payments(request):
	payments = md.Payment.objects.filter(is_deleted=False)
	template="payments.html"
	context = {'payments' : payments}
	return render(request,template , context)

def products(request):
	products = md.Product.objects.all()
	template="products.html"
	context = {'products' : products, 'firm': get_firm()}
	return render(request,template , context)

def sizes(request):
	if request.method == 'POST' :
		size_form = md.SizeForm(request.POST)
		if size_form.is_valid():
			temp_size = size_form.save(commit=True)
			md.log_audit_entry("create", temp_size, actor=request.user, after_data=md.audit_snapshot(temp_size))
			return HttpResponseRedirect('/sizes/')
	else:
		size_form = md.SizeForm()
	sizes = md.Size.objects.all()
	template = 'sizes.html'
	context = {'sizes' : sizes , 'form' : size_form}
	return render(request,template , context)	


@transaction.atomic
def delete_size(request, id):
	if request.method != "POST":
		return HttpResponseRedirect("/sizes/")
	if not get_firm().show_size_delete_button:
		messages.error(request, "Size deletion is disabled in Firm Settings.")
		return HttpResponseRedirect("/sizes/")
	try:
		size = md.Size.objects.get(id=id)
	except md.Size.DoesNotExist:
		return HttpResponseRedirect("/sizes/")
	if md.Item.objects.filter(size=size).exists():
		messages.error(request, "This size is already used in invoice items and cannot be deleted.")
		return HttpResponseRedirect("/sizes/")
	before = md.audit_snapshot(size)
	size_name = size.name
	size_id = size.id
	size.delete()
	md.log_audit_entry(
		"delete",
		size,
		actor=request.user,
		before_data=before,
		after_data={"deleted": True, "size_id": size_id, "name": size_name},
	)
	messages.success(request, "Size deleted successfully.")
	return HttpResponseRedirect("/sizes/")


def firm_settings(request):
	firm = md.FirmSettings.current()
	if request.method == 'POST':
		before = md.audit_snapshot(firm)
		form = md.FirmSettingsForm(request.POST, request.FILES, instance=firm)
		if form.is_valid():
			updated_firm = form.save()
			md.log_audit_entry("update", updated_firm, actor=request.user, before_data=before, after_data=md.audit_snapshot(updated_firm))
			return HttpResponseRedirect("/firm-settings/")
	else:
		form = md.FirmSettingsForm(instance=firm)
	return render(request, "firm_settings.html", {"form": form, "firm": firm})


def reports(request):
	context = _build_reports_payload()
	template = "reports.html"
	return render(request, template, context)

def reports_export(request, export_type):
	payload = _build_reports_payload()

	if export_type == "xlsx":
		wb = Workbook()
		summary_ws = wb.active
		summary_ws.title = "Summary"
		summary_ws.append(["Metric", "Value"])
		summary_rows = [
			("Period Start", payload["start_date"]),
			("Period End", payload["end_date"]),
			("Sales", payload["sales_total"]),
			("Collections", payload["collections_total"]),
			("Receivables", payload["total_receivables"]),
			("Pending Invoices", payload["pending_invoice_count"]),
			("Sales", payload["sales_total"]),
			("CGST", payload["gst_summary"]["cgst_total"]),
			("SGST", payload["gst_summary"]["sgst_total"]),
			("IGST", payload["gst_summary"]["igst_total"]),
			("Total Tax", payload["gst_summary"]["total_tax"]),
		]
		for item in summary_rows:
			summary_ws.append([_format_export_value(item[0]), _format_export_value(item[1])])
		_style_workbook_sheet(summary_ws)

		sheets = [
			("Monthly Movement", ["Month", "Sales", "Collections", "Net Movement"], [
				[row["month"], row["sales"], row["collections"], row["net_movement"]] for row in payload["monthly_rows"]
			]),
			("GST Summary", ["Metric", "Value"], [
				["Sales", payload["sales_total"]],
				["Base Amount", payload["gst_summary"]["base_total"]],
				["CGST", payload["gst_summary"]["cgst_total"]],
				["SGST", payload["gst_summary"]["sgst_total"]],
				["IGST", payload["gst_summary"]["igst_total"]],
				["Total Tax", payload["gst_summary"]["total_tax"]],
			]),
			("Receivables", ["Client", "Opening", "Invoice", "Payment", "Closing"], [
				[balance["name"], balance["opening_balance"], balance["invoice_total"], balance["payment_total"], balance["closing_balance"]]
				for balance in payload["receivables"]
			]),
			("Aging Summary", ["Bucket", "Count", "Amount"], [
				[bucket["label"], bucket["count"], bucket["amount"]] for bucket in payload["aging_buckets"]
			]),
			("Agents", ["Agent", "Invoices", "Value"], [
				[row["agent__name"] or "-", row["invoice_count"], row["total"]] for row in payload["agent_rows"]
			]),
			("Products", ["Product", "Qty", "Total"], [
				[row["product__name"] or "-", row["qty_total"], row["total"]] for row in payload["product_rows"]
			]),
			("Recent Activity", ["Date", "Type", "Party", "Title", "Status", "Amount", "URL"], [
				[row["date"], row["kind"], row["party"], row["title"], row["status"], row["amount"], row["url"]]
				for row in payload["recent_activity"]
			]),
		]
		for sheet_name, headers, rows in sheets:
			ws = wb.create_sheet(title=sheet_name)
			_append_sheet_rows(ws, headers, rows)

		buffer = io.BytesIO()
		wb.save(buffer)
		response = HttpResponse(
			buffer.getvalue(),
			content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
		)
		response["Content-Disposition"] = f'attachment; filename="{_reports_filename("xlsx")}"'
		return response

	if export_type == "csv":
		zip_buffer = io.BytesIO()
		with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as archive:
			csv_sheets = [
				("summary.csv", ["Metric", "Value"], [
					["Period Start", payload["start_date"]],
					["Period End", payload["end_date"]],
					["Sales", payload["sales_total"]],
					["Collections", payload["collections_total"]],
					["Receivables", payload["total_receivables"]],
					["Pending Invoices", payload["pending_invoice_count"]],
				]),
				("monthly_movement.csv", ["Month", "Sales", "Collections", "Net Movement"], [
					[row["month"], row["sales"], row["collections"], row["net_movement"]] for row in payload["monthly_rows"]
				]),
				("gst_summary.csv", ["Metric", "Value"], [
					["Sales", payload["sales_total"]],
					["Base Amount", payload["gst_summary"]["base_total"]],
					["CGST", payload["gst_summary"]["cgst_total"]],
					["SGST", payload["gst_summary"]["sgst_total"]],
					["IGST", payload["gst_summary"]["igst_total"]],
					["Total Tax", payload["gst_summary"]["total_tax"]],
				]),
				("receivables.csv", ["Client", "Opening", "Invoice", "Payment", "Closing"], [
					[balance["name"], balance["opening_balance"], balance["invoice_total"], balance["payment_total"], balance["closing_balance"]]
					for balance in payload["receivables"]
				]),
				("aging_summary.csv", ["Bucket", "Count", "Amount"], [
					[bucket["label"], bucket["count"], bucket["amount"]] for bucket in payload["aging_buckets"]
				]),
				("agents.csv", ["Agent", "Invoices", "Value"], [
					[row["agent__name"] or "-", row["invoice_count"], row["total"]] for row in payload["agent_rows"]
				]),
				("products.csv", ["Product", "Qty", "Total"], [
					[row["product__name"] or "-", row["qty_total"], row["total"]] for row in payload["product_rows"]
				]),
				("recent_activity.csv", ["Date", "Type", "Party", "Title", "Status", "Amount", "URL"], [
					[row["date"], row["kind"], row["party"], row["title"], row["status"], row["amount"], row["url"]]
					for row in payload["recent_activity"]
				]),
			]
			for filename, headers, rows in csv_sheets:
				archive.writestr(filename, _csv_bytes(headers, rows))
		response = HttpResponse(zip_buffer.getvalue(), content_type="application/zip")
		response["Content-Disposition"] = f'attachment; filename="{_reports_filename("zip")}"'
		return response

	return HttpResponseRedirect("/reports/")

def maintenance(request):
	backup_folder = _backup_folder()
	latest_snapshot = sorted(backup_folder.glob("*.zip"), key=lambda path: path.stat().st_mtime, reverse=True)
	current_backup_size = Path(settings.DATABASES["default"]["NAME"]).stat().st_size if Path(settings.DATABASES["default"]["NAME"]).exists() else 0
	if request.method == "POST":
		if request.POST.get("action") == "restore":
			uploaded_file = request.FILES.get("backup_file")
			confirm_restore = request.POST.get("confirm_restore") == "on"
			if not uploaded_file:
				messages.error(request, "Please choose a backup ZIP file to restore.")
			elif not confirm_restore:
				messages.error(request, "Please confirm that you want to restore this backup.")
			else:
				try:
					snapshot_path = _restore_from_backup_file(uploaded_file)
					messages.success(request, f"Backup restored successfully. Safety snapshot saved at {snapshot_path}.")
					return HttpResponseRedirect("/maintenance/")
				except Exception as exc:
					messages.error(request, f"Restore failed: {exc}")
		elif request.POST.get("action") == "truncate_database":
			confirm_truncate = request.POST.get("confirm_truncate") == "on"
			if not confirm_truncate:
				messages.error(request, "Please confirm that you want to remove invoice, payment, and product data.")
			else:
				try:
					snapshot_path = _truncate_company_database()
					messages.success(
						request,
						f"Invoice, payment, and product data were removed successfully. Safety snapshot saved at {snapshot_path}.",
					)
					return HttpResponseRedirect("/maintenance/")
				except Exception as exc:
					messages.error(request, f"Database wipe failed: {exc}")
		elif request.POST.get("action") == "create_backup":
			try:
				snapshot_path = _save_backup_snapshot("manual")
				messages.success(request, f"Backup created at {snapshot_path}.")
				return HttpResponseRedirect(f"/maintenance/?download={snapshot_path.name}")
			except Exception as exc:
				messages.error(request, f"Backup creation failed: {exc}")

	download_name = request.GET.get("download")
	if download_name:
		candidate = backup_folder / download_name
		if candidate.exists():
			with open(candidate, "rb") as handle:
				response = HttpResponse(handle.read(), content_type="application/zip")
			response["Content-Disposition"] = f'attachment; filename="{candidate.name}"'
			return response

	context = {
		"backup_folder": backup_folder,
		"latest_snapshot": latest_snapshot[0] if latest_snapshot else None,
		"db_path": Path(settings.DATABASES["default"]["NAME"]),
		"db_size": current_backup_size,
	}
	return render(request, "maintenance.html", context)

def download_backup(request):
	response = HttpResponse(_build_backup_archive().getvalue(), content_type="application/zip")
	filename = f"accountmanager-backup-{timezone.localtime().strftime('%Y%m%d_%H%M%S')}.zip"
	response["Content-Disposition"] = f'attachment; filename="{filename}"'
	return response


#####################################################################################################################
##			ADD SOMETHING
#####################################################################################################################
def new_agent(request):
	if request.method == 'POST':
		temp_form = md.AgentForm(request.POST)
		if temp_form.is_valid():
			temp_agent = temp_form.save(commit=False)
			temp_agent.save()
			md.log_audit_entry("create", temp_agent, actor=request.user, after_data=md.audit_snapshot(temp_agent))
			return HttpResponseRedirect("/agents")

	temp_form = md.AgentForm()
	template = "new_agent.html"
	context = {'form' : temp_form}
	return render(request , template, context)

def new_client(request):
	next_url = request.GET.get("next") or request.POST.get("next")
	if request.method == 'POST' : 
		temp_form = md.ClientForm(request.POST)
		if temp_form.is_valid():
			temp_client = temp_form.save(commit=False)
			temp_client.debit_balance = 0.0
			temp_client.save()
			md.log_audit_entry("create", temp_client, actor=request.user, after_data=md.audit_snapshot(temp_client))
			redirect_to = _redirect_with_client(next_url, temp_client.id)
			if redirect_to:
				return HttpResponseRedirect(redirect_to)
			return HttpResponseRedirect("/invoices/new")
		else:
			temp_form = md.ClientForm(request.POST)
	else:
		temp_form = md.ClientForm()
	template="new_client.html"
	context = {"form" : temp_form , "states" : md.STATES, "next_url": next_url}
	return render(request,template , context)

def new_invoice(request):
	errors = None
	formset = formset_factory(md.ItemForm , extra=1)
	selected_client = request.GET.get("client")
	if request.method == 'POST' :
		## Add a dynamic Agent name if it doesn't exists
		request.POST._mutable=True
		request.POST['agent'] = process_agent(request.POST.get('agent', ''))
		request.POST._mutable=False
		
		form = md.BillForm(request.POST)
		Itemformset = formset(request.POST)
		if form.is_valid() and Itemformset.is_valid():
			with transaction.atomic():
				bill = form.save(commit=False)
				bill.is_taxable = True
				
				parcel_digits = request.POST.get("parcel_digits" , "")
				bill.invoice_number = get_invoice_number(parcel_digits)
				bill.qty_total = 0
				bill.save()
				
				bill.total_before_tax = Decimal(0)
				for Itemform in Itemformset:
					item = Itemform.save(commit=False)
					item.bill= bill
					item.save()
					bill.total_before_tax += item.rate*item.qty

				bill.total_before_tax *= (100-bill.discount)*Decimal(0.01) 
				_apply_bill_tax_rates(bill)
				if bill.client.state_code == '24':
					tax_multiplier = Decimal("1.00") + ((bill.cgst + bill.sgst) / Decimal("100.00"))
				else:
					tax_multiplier = Decimal("1.00") + (bill.igst / Decimal("100.00"))
				bill.total_after_tax = Decimal(math.ceil(float(bill.total_before_tax) * float(tax_multiplier)))

				bill.qty_total_set()
				bill.save()
				_refresh_client_balance(bill.client)
				md.log_audit_entry(
					"create",
					bill,
					actor=request.user,
					after_data=md.audit_snapshot(bill, include_related=True),
				)
			return HttpResponseRedirect("/invoices/{0}".format(bill.id))
		else:
			print(form.errors)
			print(Itemformset.errors)
	else:	
		form = md.BillForm()
	template="new_invoice.html"
	context = {
		"form" : form,
		"formset" : formset(),
		"errors" : errors,
		"size_choices" : [x.name for x in md.Size.objects.all()],
		"product_choices" : [x.name for x in md.Product.objects.all()],
		"add" : True,
		"name" : "New Invoice",
		"selected_client": selected_client,
	}
	return render(request,template , context)

def new_payment(request):
	if request.method == 'POST' : 
		temp_form = md.PaymentForm(request.POST)
		if temp_form.is_valid():
			with transaction.atomic():
				temp_payment = temp_form.save(commit=False)
				temp_payment.save()
				_refresh_client_balance(temp_payment.client)
				md.log_audit_entry("create", temp_payment, actor=request.user, after_data=md.audit_snapshot(temp_payment))
			return HttpResponseRedirect("/payments")
		else:
			return HttpResponse(str(temp_form.errors))
	temp_form = md.PaymentForm()
	template="new_payment.html"
	context = {"form" : temp_form , "add" : True}
	return render(request,template , context)

def new_product(request):
	if request.method == 'POST' : 
		temp_form = md.ProductForm(request.POST)
		if temp_form.is_valid():
			temp_product = temp_form.save(commit=False)
			temp_product.save()
			md.log_audit_entry("create", temp_product, actor=request.user, after_data=md.audit_snapshot(temp_product))
			return HttpResponseRedirect("/products")
		else:
			temp_form = md.ProductForm(request.POST)
	else:
		temp_form = md.ProductForm()
	template="new_product.html"
	context = {'form' : temp_form}
	return render(request,template , context)	


#####################################################################################################################
##			Individual Page and Edit
#####################################################################################################################
def single_invoice(request , id):
	try:
		invoice = md.Bill.objects.get(id=id, is_deleted=False)
	except:
		return HttpResponseRedirect("/invoices")

	firm = get_firm()
	context={'invoice': invoice, "firm": firm, "firm_contact_numbers": _firm_contact_numbers(firm), "invoice_terms_lines": _invoice_terms_lines(firm), "auto_print": False}
	context["invoice_pages"] = _invoice_pages(invoice)
	
	template = _invoice_template_name()
	return render(request,template , context)


def print_invoice(request, id):
	try:
		invoice = md.Bill.objects.get(id=id, is_deleted=False)
	except:
		return HttpResponseRedirect("/invoices")

	firm = get_firm()
	context = {'invoice': invoice, "firm": firm, "firm_contact_numbers": _firm_contact_numbers(firm), "invoice_terms_lines": _invoice_terms_lines(firm), "auto_print": True}
	context["invoice_pages"] = _invoice_pages(invoice)
	return render(request, _print_invoice_template_name(), context)

def single_product(request,id):
	try:
		product = md.Product.objects.get(id=id)
	except:
		return HttpResponseRedirect("/products")
	
	template = 'single_product.html'
	context = {'product' : product}
	return render(request,template , context)


@transaction.atomic
def delete_product(request, id):
	if request.method != "POST":
		return HttpResponseRedirect("/products/")
	if not get_firm().show_product_delete_button:
		messages.error(request, "Product deletion is disabled in Firm Settings.")
		return HttpResponseRedirect("/products/")
	try:
		product = md.Product.objects.get(id=id)
	except md.Product.DoesNotExist:
		return HttpResponseRedirect("/products")
	before = md.audit_snapshot(product)
	product_name = product.name
	product_id = product.id
	affected_bills = list(
		md.Bill.objects.filter(item__product=product, is_deleted=False).distinct().select_related("client")
	)
	product.delete()
	for bill in affected_bills:
		_refresh_bill_totals(bill)
		_refresh_client_balance(bill.client)
	md.log_audit_entry(
		"delete",
		product,
		actor=request.user,
		before_data=before,
		after_data={
			"deleted": True,
			"product_id": product_id,
			"name": product_name,
			"affected_bill_ids": [bill.id for bill in affected_bills],
		},
	)
	messages.success(request, "Product deleted successfully.")
	return HttpResponseRedirect("/products")

def single_agent(request,id):
	try:
		agent = md.Agent.objects.get(id=id)
	except:
		return HttpResponseRedirect("/agents")
	
	template = 'single_agent.html'
	context = {'agent' : agent}
	return render(request,template , context)


@transaction.atomic
def delete_agent(request, id):
	if request.method != "POST":
		return HttpResponseRedirect("/agents/")
	if not get_firm().show_agent_delete_button:
		messages.error(request, "Agent deletion is disabled in Firm Settings.")
		return HttpResponseRedirect("/agents/")
	try:
		agent = md.Agent.objects.get(id=id)
	except md.Agent.DoesNotExist:
		return HttpResponseRedirect("/agents")
	before = md.audit_snapshot(agent)
	agent_name = agent.name
	agent_id = agent.id
	affected_bill_ids = list(md.Bill.objects.filter(agent=agent, is_deleted=False).values_list("id", flat=True))
	md.Bill.objects.filter(agent=agent).update(agent=None)
	agent.delete()
	md.log_audit_entry(
		"delete",
		agent,
		actor=request.user,
		before_data=before,
		after_data={
			"deleted": True,
			"agent_id": agent_id,
			"name": agent_name,
			"affected_bill_ids": affected_bill_ids,
		},
	)
	messages.success(request, "Agent deleted successfully.")
	return HttpResponseRedirect("/agents")

def single_client(request,id):
	try: client = md.Client.objects.get(id=id)
	except: return HttpResponseRedirect("/clients")
	client.refresh_debit_balance()
	
	template = 'single_client.html'
	context = {'client' : client}
	return render(request,template , context)


@transaction.atomic
def delete_client(request, id):
	if request.method != "POST":
		return HttpResponseRedirect("/clients/{0}".format(id))
	if not get_firm().show_client_delete_button:
		messages.error(request, "Client deletion is disabled in Firm Settings.")
		return HttpResponseRedirect("/clients/{0}".format(id))
	try:
		client = md.Client.objects.get(id=id)
	except md.Client.DoesNotExist:
		return HttpResponseRedirect("/clients")
	before = md.audit_snapshot(client)
	client_name = client.name
	client_id = client.id
	client.delete()
	md.log_audit_entry(
		"delete",
		client,
		actor=request.user,
		before_data=before,
		after_data={"deleted": True, "client_id": client_id, "name": client_name},
	)
	messages.success(request, "Client deleted successfully.")
	return HttpResponseRedirect("/clients")

def single_payment(request,id):
	try: payment = md.Payment.objects.get(id=id, is_deleted=False)
	except: return HttpResponseRedirect("/payments")
	if payment.is_deleted:
		return HttpResponseRedirect("/payments")
	
	template = 'single_payment.html'
	context = {'payment' : payment}
	return render(request,template , context)	


@transaction.atomic
def delete_invoice(request, id):
	if request.method != "POST":
		return HttpResponseRedirect("/invoices/{0}".format(id))
	try:
		bill = md.Bill.objects.select_related("client").get(id=id, is_deleted=False)
	except md.Bill.DoesNotExist:
		return HttpResponseRedirect("/invoices")
	before = md.audit_snapshot(bill, include_related=True)
	client = bill.client
	bill.is_deleted = True
	bill.is_cleared = False
	bill.save(update_fields=["is_deleted", "is_cleared"])
	_refresh_client_balance(client)
	md.log_audit_entry("delete", bill, actor=request.user, before_data=before, after_data=md.audit_snapshot(bill, include_related=True))
	messages.success(request, "Invoice deleted successfully.")
	return HttpResponseRedirect("/invoices")


@transaction.atomic
def delete_payment(request, id):
	if request.method != "POST":
		return HttpResponseRedirect("/payments/{0}".format(id))
	if not get_firm().show_payment_delete_button:
		messages.error(request, "Payment deletion is disabled in Firm Settings.")
		return HttpResponseRedirect("/payments/{0}".format(id))
	try:
		payment = md.Payment.objects.select_related("client").get(id=id, is_deleted=False)
	except md.Payment.DoesNotExist:
		return HttpResponseRedirect("/payments")
	before = md.audit_snapshot(payment)
	client = payment.client
	payment.is_deleted = True
	payment.is_cleared = False
	payment.save(update_fields=["is_deleted", "is_cleared"])
	_refresh_client_balance(client)
	md.log_audit_entry("delete", payment, actor=request.user, before_data=before, after_data=md.audit_snapshot(payment))
	messages.success(request, "Payment deleted successfully.")
	return HttpResponseRedirect("/payments")

def edit_payment(request,id):
	try : 
		payment = md.Payment.objects.get(id=id, is_deleted=False)
		old_client = payment.client
		old_amount = payment.paid_amount
		before = md.audit_snapshot(payment)
	except: return HttpResponseRedirect("/payments/")
	if request.method == 'POST' : 
		temp_form = md.PaymentForm(request.POST , instance=payment)
		if temp_form.is_valid():
			temp_payment = temp_form.save(commit=False)
			temp_payment.save()
			_refresh_client_balance(old_client)
			_refresh_client_balance(temp_payment.client)
			md.log_audit_entry("update", temp_payment, actor=request.user, before_data=before, after_data=md.audit_snapshot(temp_payment))
			return HttpResponseRedirect("/payments/{0}".format(temp_payment.id))
		else:
			return HttpResponse(str(temp_form.errors))
	temp_form = md.PaymentForm(instance=payment)
	template="new_payment.html"
	context = {"form" : temp_form}
	return render(request,template , context)

def edit_client(request,id):
	try: client = md.Client.objects.get(id=id)
	except: return HttpResponseRedirect("/clients")
	if request.method == 'POST' :
		before = md.audit_snapshot(client)
		temp_form = md.ClientForm(request.POST , instance=client)
		if temp_form.is_valid():
			temp_client = temp_form.save(commit=False)
			temp_client.save()
			md.log_audit_entry("update", temp_client, actor=request.user, before_data=before, after_data=md.audit_snapshot(temp_client))
			return HttpResponseRedirect("/clients/{0}".format(client.id))
		else:
			temp_form = md.ClientForm(request.POST)
	
	temp_form = md.ClientForm(instance=client)
	template="new_client.html"
	context={"form" : temp_form}
	return render(request,template , context)

def edit_invoice(request,id):
	error = None
	try:
		old_bill = md.Bill.objects.get(id=id, is_deleted=False)
		old_items = md.Item.objects.filter(bill=old_bill)
		old_items_id = [item.id for item in old_items]
		old_bill_value = old_bill.total_after_tax
		old_bill_client = old_bill.client
		before = md.audit_snapshot(old_bill, include_related=True)
	except:
		return HttpResponseRedirect("/invoices")
	formset = modelformset_factory(md.Item, form=md.ItemForm, extra=0, can_delete=True, max_num=100 )
	if (request.method == 'POST'):
		## Add a dynamic Agent name if it doesn't exists
		request.POST._mutable=True
		request.POST['agent'] = process_agent(request.POST.get('agent', ''))
		request.POST._mutable=False

		form = md.BillForm(request.POST, instance=old_bill)
		Itemformset = formset(request.POST , queryset=old_items)
		if form.is_valid() and Itemformset.is_valid():
			with transaction.atomic():
				bill = form.save(commit=False)
				bill.is_taxable = True
				parcel_digits = request.POST.get("parcel_digits" , "")
				invoice_parts = old_bill.invoice_number.split("/")
				if len(invoice_parts) >= 3 and "-" in invoice_parts[1]:
					invoice_base = "/".join(invoice_parts[:3])
				else:
					invoice_base = "/".join(invoice_parts[:2])
				if parcel_digits and parcel_digits.strip().isalnum():
					bill.invoice_number = invoice_base + "/{0}".format(parcel_digits.strip())
				else:
					bill.invoice_number = old_bill.invoice_number
				bill.qty_total = 0
				bill.save()
				
				bill.total_before_tax = Decimal(0)
				for Itemform in Itemformset:
					item = Itemform.save(commit=False)
					if item.id in old_items_id: old_items_id.remove(item.id)
					item.bill= bill
					item.save()
					bill.total_before_tax += item.rate*item.qty
				for item in old_items.filter(id__in=old_items_id):
					item.delete()
				
				bill.qty_total_set()
				bill.total_before_tax *= (100-bill.discount)*Decimal(0.01) 
				_apply_bill_tax_rates(bill)
				if bill.client.state_code == '24':
					tax_multiplier = Decimal("1.00") + ((bill.cgst + bill.sgst) / Decimal("100.00"))
				else:
					tax_multiplier = Decimal("1.00") + (bill.igst / Decimal("100.00"))
				bill.total_after_tax = Decimal(math.ceil(float(bill.total_before_tax) * float(tax_multiplier)))
				bill.save()
				_refresh_client_balance(old_bill.client)
				_refresh_client_balance(bill.client)
				md.log_audit_entry("update", bill, actor=request.user, before_data=before, after_data=md.audit_snapshot(bill, include_related=True))
			return HttpResponseRedirect("/invoices/{0}".format(bill.id))
		else: 
			print(form.errors)
			fm = formset()
			print(fm.non_form_errors() )
	else:	
		form = md.BillForm(instance=old_bill)
	template="new_invoice.html"
	context = {
		"form" : form,
		"formset" : formset(queryset = old_items),
		"name" : "Edit Invoice",
		"size_choices" : [x.name for x in md.Size.objects.all()],
		"product_choices" : [x.name for x in md.Product.objects.all()],
		"selected_client": old_bill.client.id,
	}
	return render(request,template , context)
